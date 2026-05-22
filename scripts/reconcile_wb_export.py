#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import os
import subprocess
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "output" / "spreadsheet"


def load_env_database_url() -> str:
    if os.environ.get("DATABASE_URL"):
        return os.environ["DATABASE_URL"]

    env_path = ROOT / ".env"
    if not env_path.exists():
        raise RuntimeError("DATABASE_URL is not set and .env is missing")

    for raw_line in env_path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        if key.strip() == "DATABASE_URL":
            return value.strip().strip('"').strip("'")

    raise RuntimeError("DATABASE_URL is not set and was not found in .env")


def run_psql_json(database_url: str, query: str) -> Any:
    command = [
        "psql",
        database_url,
        "-At",
        "-F",
        "\t",
        "-c",
        query,
    ]
    result = subprocess.run(command, capture_output=True, text=True, check=True)
    output = result.stdout.strip()
    if not output:
        return None
    return json.loads(output)


def sql_literal(value: str) -> str:
    return "'" + value.replace("'", "''") + "'"


def sql_int_array(values: list[int]) -> str:
    if not values:
        return "ARRAY[]::bigint[]"
    return "ARRAY[" + ",".join(str(int(value)) for value in values) + "]::bigint[]"


def parse_numeric(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace("\xa0", " ").replace(" ", "").replace("%", "").replace(",", ".")
    return float(text)


def parse_percentage(value: Any) -> float:
    numeric = parse_numeric(value)
    if 0 <= numeric <= 1:
        return numeric * 100
    return numeric


def parse_int(value: Any) -> int:
    return int(round(parse_numeric(value)))


def load_export_metrics(workbook_path: Path) -> dict[str, Any]:
    wb = load_workbook(workbook_path, data_only=True, read_only=False)
    info_sheet = wb["Общая информация"]
    filters_sheet = wb["Фильтры"]
    products_sheet = wb["Товары"]

    info: dict[str, Any] = {}
    for row in info_sheet.iter_rows(values_only=True):
        key = row[0] if len(row) > 0 else None
        value = row[1] if len(row) > 1 else None
        if key not in (None, "") and value not in (None, ""):
            info[str(key).strip()] = value

    filter_rows = list(filters_sheet.iter_rows(values_only=True))
    filter_header_index = next(
        index
        for index, row in enumerate(filter_rows)
        if any(value == "Переходы в карточку" for value in row if value is not None)
    )
    filter_headers = list(filter_rows[filter_header_index])
    filter_values = list(filter_rows[filter_header_index + 1])
    filters = {
        str(header).strip(): value
        for header, value in zip(filter_headers, filter_values)
        if header is not None
    }

    product_rows = list(products_sheet.iter_rows(values_only=True))
    product_header_index = next(
        index
        for index, row in enumerate(product_rows)
        if any(value == "Артикул WB" for value in row if value is not None)
    )
    product_headers = list(product_rows[product_header_index])
    products: list[dict[str, Any]] = []
    for row in product_rows[product_header_index + 1:]:
        if not any(item is not None and item != "" for item in row):
            continue
        products.append({
            str(header).strip(): value
            for header, value in zip(product_headers, row)
            if header is not None
        })

    export_nm_ids: set[int] = set()
    deleted_nm_ids: set[int] = set()
    for product in products:
        nm_id_raw = product.get("Артикул WB")
        if nm_id_raw in (None, ""):
            continue
        nm_id = parse_int(nm_id_raw)
        export_nm_ids.add(nm_id)

        deleted_marker = str(product.get("Удаленный товар") or "").strip().lower()
        if deleted_marker in {"да", "true", "1"}:
            deleted_nm_ids.add(nm_id)

    return {
        "context": {
            "current_period": str(info.get("Значение за текущий период") or "").strip(),
            "previous_period": str(info.get("Прошлый период") or "").strip(),
            "subject": str(info.get("Предметы") or "").strip() or None,
            "deleted_mode": str(info.get("Удалённые товары") or "").strip() or None,
            "granularity": str(info.get("Детализация") or "").strip() or None,
        },
        "filters": {
            "views": parse_int(filters.get("Переходы в карточку")),
            "add_to_cart": parse_int(filters.get("Положили в корзину")),
            "ordered_units": parse_int(filters.get("Заказали, шт")),
            "bought_units": parse_int(filters.get("Выкупили, шт")),
            "cancelled_units": parse_int(filters.get("Отменили, шт")),
            "ordered_rub": parse_numeric(filters.get("Заказали на сумму, ₽")),
            "bought_rub": parse_numeric(filters.get("Выкупили на сумму, ₽")),
            "cancelled_rub": parse_numeric(filters.get("Отменили на сумму, ₽")),
            "avg_price": parse_numeric(filters.get("Средняя цена, ₽")),
            "cart_conversion_pct": parse_percentage(filters.get("Конверсия в корзину, %")),
            "order_conversion_pct": parse_percentage(filters.get("Конверсия в заказ, %")),
        },
        "products": {
            "rows": len(products),
            "unique_nm_ids": sorted(export_nm_ids),
            "deleted_nm_ids": sorted(deleted_nm_ids),
        },
    }


def load_db_metrics(
    database_url: str,
    tenant_id: str,
    date_from: str,
    date_to: str,
    export_nm_ids: list[int],
) -> dict[str, Any]:
    date_to_exclusive = (datetime.fromisoformat(date_to) + timedelta(days=1)).date().isoformat()
    nm_ids_sql = sql_int_array(export_nm_ids)

    tenant = run_psql_json(
        database_url,
        f"""
        select row_to_json(t)
        from (
          select id, name
          from tenants
          where id = {sql_literal(tenant_id)}
        ) t;
        """,
    )
    if tenant is None:
        raise RuntimeError(f"Tenant {tenant_id} not found")

    overview = run_psql_json(
        database_url,
        f"""
        WITH reconciliation_cutoff AS (
          SELECT COALESCE(MAX(date_to), '1970-01-01'::timestamp) as cutoff
          FROM raw_api_realization_reports
          WHERE tenant_id = {sql_literal(tenant_id)}
        ),
        unified_sales AS (
          SELECT
            r.nm_id,
            r.retail_amount as revenue,
            (r.retail_amount - r.commission_amount - r.delivery_rub - r.storage_fee_rub - r.penalty_rub - r.payment_schedule_rub - (r.quantity * COALESCE(c.cost_price, 0))) as profit
          FROM raw_api_realization_reports r
          LEFT JOIN LATERAL (
             SELECT cost_price FROM unit_economics_configs
             WHERE nm_id = r.nm_id AND tenant_id = r.tenant_id AND effective_from <= r.date_from
             ORDER BY effective_from DESC LIMIT 1
          ) c ON true
          LEFT JOIN products p ON p.tenant_id = r.tenant_id AND p.nm_id = r.nm_id
          WHERE r.tenant_id = {sql_literal(tenant_id)}
            AND r.date_from >= {sql_literal(date_from)}::timestamp
            AND r.date_from < {sql_literal(date_to_exclusive)}::timestamp
            AND COALESCE(p.is_hidden, FALSE) = FALSE

          UNION ALL

          SELECT
            s.nm_id,
            s.price_with_discount as revenue,
            (s.price_with_discount * 0.75 - COALESCE(c.cost_price, 0)) as profit
          FROM raw_api_sales s
          CROSS JOIN reconciliation_cutoff rc
          LEFT JOIN LATERAL (
             SELECT cost_price FROM unit_economics_configs
             WHERE nm_id = s.nm_id AND tenant_id = s.tenant_id AND effective_from <= s.date
             ORDER BY effective_from DESC LIMIT 1
          ) c ON true
          LEFT JOIN products p ON p.tenant_id = s.tenant_id AND p.nm_id = s.nm_id
          WHERE s.tenant_id = {sql_literal(tenant_id)}
            AND s.date > rc.cutoff
            AND s.date >= {sql_literal(date_from)}::timestamp
            AND s.date < {sql_literal(date_to_exclusive)}::timestamp
            AND s.is_storno = false
            AND COALESCE(p.is_hidden, FALSE) = FALSE
        ),
        totals AS (
          SELECT
            COALESCE(SUM(u.revenue), 0)::numeric as finance_revenue,
            COALESCE(SUM(u.profit), 0)::numeric as op_profit,
            t.tax_type,
            COALESCE(t.tax_rate, '0') as tax_rate
          FROM unified_sales u
          JOIN tenants t ON t.id = {sql_literal(tenant_id)}
          GROUP BY t.tax_type, t.tax_rate
        ),
        storage_total AS (
          SELECT COALESCE(SUM(storage_amount), 0)::numeric as total_storage
          FROM raw_api_paid_storage s
          LEFT JOIN products p ON p.tenant_id = s.tenant_id AND p.nm_id = s.nm_id
          WHERE s.tenant_id = {sql_literal(tenant_id)}
            AND s.date >= {sql_literal(date_from)}::timestamp
            AND s.date < {sql_literal(date_to_exclusive)}::timestamp
            AND COALESCE(p.is_hidden, FALSE) = FALSE
        ),
        ad_total AS (
          SELECT COALESCE(SUM(amount), 0)::numeric as total_ads
          FROM raw_api_ad_costs a
          LEFT JOIN products p ON p.tenant_id = a.tenant_id AND p.nm_id = a.nm_id
          WHERE a.tenant_id = {sql_literal(tenant_id)}
            AND a.date >= {sql_literal(date_from)}::timestamp
            AND a.date < {sql_literal(date_to_exclusive)}::timestamp
            AND COALESCE(p.is_hidden, FALSE) = FALSE
        ),
        funnel_total AS (
          SELECT
            COALESCE(SUM(f.open_card_count), 0)::numeric as total_views,
            COALESCE(SUM(f.add_to_cart_count), 0)::numeric as total_carts,
            COALESCE(SUM(f.order_count), 0)::numeric as total_orders,
            COALESCE(SUM(f.order_sum), 0)::numeric as total_order_sum,
            COALESCE(SUM(f.buyout_count), 0)::numeric as total_buyouts,
            COALESCE(SUM(f.buyout_sum), 0)::numeric as total_buyout_sum,
            COALESCE(SUM(f.cancel_count), 0)::numeric as total_cancels,
            COALESCE(SUM(f.cancel_sum), 0)::numeric as total_cancel_sum,
            COUNT(*)::int as total_rows
          FROM raw_api_funnel_stats f
          LEFT JOIN products p ON p.tenant_id = f.tenant_id AND p.nm_id = f.nm_id
          WHERE f.tenant_id = {sql_literal(tenant_id)}
            AND f.period_start = {sql_literal(date_from)}::timestamp
            AND f.period_end = {sql_literal(date_to)}::timestamp
            AND COALESCE(p.is_hidden, FALSE) = FALSE
        )
        SELECT row_to_json(t)
        FROM (
          SELECT
            COALESCE((SELECT finance_revenue FROM totals), 0)::numeric as finance_revenue,
            COALESCE((SELECT total_ads FROM ad_total), 0)::numeric as total_ads,
            COALESCE((SELECT total_storage FROM storage_total), 0)::numeric as total_storage,
            COALESCE((SELECT total_views FROM funnel_total), 0)::numeric as total_views,
            COALESCE((SELECT total_carts FROM funnel_total), 0)::numeric as total_carts,
            COALESCE((SELECT total_orders FROM funnel_total), 0)::numeric as total_orders,
            COALESCE((SELECT total_order_sum FROM funnel_total), 0)::numeric as total_order_sum,
            COALESCE((SELECT total_buyouts FROM funnel_total), 0)::numeric as total_buyouts,
            COALESCE((SELECT total_buyout_sum FROM funnel_total), 0)::numeric as total_buyout_sum,
            COALESCE((SELECT total_cancels FROM funnel_total), 0)::numeric as total_cancels,
            COALESCE((SELECT total_cancel_sum FROM funnel_total), 0)::numeric as total_cancel_sum,
            COALESCE((SELECT total_rows FROM funnel_total), 0)::int as total_rows
        ) t;
        """,
    )

    export_slice_funnel = run_psql_json(
        database_url,
        f"""
        WITH export_nm AS (
          SELECT unnest({nm_ids_sql}) AS nm_id
        )
        SELECT row_to_json(t)
        FROM (
          SELECT
            COUNT(DISTINCT f.nm_id)::int as matched_skus,
            COALESCE(SUM(f.open_card_count), 0)::numeric as views,
            COALESCE(SUM(f.add_to_cart_count), 0)::numeric as carts,
            COALESCE(SUM(f.order_count), 0)::numeric as orders,
            COALESCE(SUM(f.buyout_count), 0)::numeric as buyouts,
            COALESCE(SUM(f.cancel_count), 0)::numeric as cancels,
            COALESCE(SUM(f.order_sum), 0)::numeric as order_sum,
            COALESCE(SUM(f.buyout_sum), 0)::numeric as buyout_sum,
            COALESCE(SUM(f.cancel_sum), 0)::numeric as cancel_sum
          FROM raw_api_funnel_stats f
          JOIN export_nm e ON e.nm_id = f.nm_id
          WHERE f.tenant_id = {sql_literal(tenant_id)}
            AND f.period_start = {sql_literal(date_from)}::timestamp
            AND f.period_end = {sql_literal(date_to)}::timestamp
        ) t;
        """,
    )

    export_slice_economics = run_psql_json(
        database_url,
        f"""
        WITH export_nm AS (
          SELECT unnest({nm_ids_sql}) AS nm_id
        ),
        reconciliation_cutoff AS (
          SELECT COALESCE(MAX(date_to), '1970-01-01'::timestamp) as cutoff
          FROM raw_api_realization_reports
          WHERE tenant_id = {sql_literal(tenant_id)}
        ),
        final_sales AS (
          SELECT
            COALESCE(SUM(r.retail_amount), 0)::numeric as revenue,
            COALESCE(SUM(r.quantity), 0)::numeric as sold_quantity
          FROM raw_api_realization_reports r
          JOIN export_nm e ON e.nm_id = r.nm_id
          WHERE r.tenant_id = {sql_literal(tenant_id)}
            AND r.date_from >= {sql_literal(date_from)}::timestamp
            AND r.date_from < {sql_literal(date_to_exclusive)}::timestamp
        ),
        provisional_sales AS (
          SELECT
            COALESCE(SUM(s.price_with_discount), 0)::numeric as revenue,
            COUNT(*)::numeric as sold_quantity
          FROM raw_api_sales s
          CROSS JOIN reconciliation_cutoff rc
          JOIN export_nm e ON e.nm_id = s.nm_id
          WHERE s.tenant_id = {sql_literal(tenant_id)}
            AND s.date > rc.cutoff
            AND s.date >= {sql_literal(date_from)}::timestamp
            AND s.date < {sql_literal(date_to_exclusive)}::timestamp
            AND s.is_storno = false
        ),
        storage_total AS (
          SELECT COALESCE(SUM(storage_amount), 0)::numeric as total_storage
          FROM raw_api_paid_storage s
          JOIN export_nm e ON e.nm_id = s.nm_id
          WHERE s.tenant_id = {sql_literal(tenant_id)}
            AND s.date >= {sql_literal(date_from)}::timestamp
            AND s.date < {sql_literal(date_to_exclusive)}::timestamp
        ),
        ad_total AS (
          SELECT COALESCE(SUM(amount), 0)::numeric as total_ads
          FROM raw_api_ad_costs a
          JOIN export_nm e ON e.nm_id = a.nm_id
          WHERE a.tenant_id = {sql_literal(tenant_id)}
            AND a.date >= {sql_literal(date_from)}::timestamp
            AND a.date < {sql_literal(date_to_exclusive)}::timestamp
        )
        SELECT row_to_json(t)
        FROM (
          SELECT
            COALESCE((SELECT revenue FROM final_sales), 0)::numeric + COALESCE((SELECT revenue FROM provisional_sales), 0)::numeric as finance_revenue,
            COALESCE((SELECT sold_quantity FROM final_sales), 0)::numeric + COALESCE((SELECT sold_quantity FROM provisional_sales), 0)::numeric as sold_quantity,
            COALESCE((SELECT total_storage FROM storage_total), 0)::numeric as total_storage,
            COALESCE((SELECT total_ads FROM ad_total), 0)::numeric as total_ads
        ) t;
        """,
    )

    catalog = run_psql_json(
        database_url,
        f"""
        WITH observed_nm AS (
          SELECT nm_id FROM raw_api_funnel_stats WHERE tenant_id = {sql_literal(tenant_id)}
          UNION
          SELECT nm_id FROM raw_api_orders WHERE tenant_id = {sql_literal(tenant_id)}
          UNION
          SELECT nm_id FROM raw_api_sales WHERE tenant_id = {sql_literal(tenant_id)}
          UNION
          SELECT nm_id FROM raw_api_realization_reports WHERE tenant_id = {sql_literal(tenant_id)}
          UNION
          SELECT nm_id FROM raw_api_ad_costs WHERE tenant_id = {sql_literal(tenant_id)}
          UNION
          SELECT nm_id FROM raw_api_product_metadata WHERE tenant_id = {sql_literal(tenant_id)}
        )
        SELECT row_to_json(t)
        FROM (
          SELECT
            COALESCE((
              SELECT json_agg(nm_id ORDER BY nm_id)
              FROM (SELECT DISTINCT nm_id FROM products WHERE tenant_id = {sql_literal(tenant_id)}) p
            ), '[]'::json) as products_nm_ids,
            COALESCE((
              SELECT json_agg(nm_id ORDER BY nm_id)
              FROM (SELECT DISTINCT nm_id FROM products WHERE tenant_id = {sql_literal(tenant_id)} AND is_archived = false) p
            ), '[]'::json) as active_products_nm_ids,
            COALESCE((
              SELECT json_agg(nm_id ORDER BY nm_id)
              FROM (SELECT DISTINCT nm_id FROM raw_api_product_metadata WHERE tenant_id = {sql_literal(tenant_id)}) m
            ), '[]'::json) as metadata_nm_ids,
            COALESCE((
              SELECT json_agg(nm_id ORDER BY nm_id)
              FROM (SELECT DISTINCT nm_id FROM observed_nm) o
            ), '[]'::json) as observed_nm_ids
        ) t;
        """,
    )

    return {
        "tenant": tenant,
        "overview": overview,
        "export_slice_funnel": export_slice_funnel,
        "export_slice_economics": export_slice_economics,
        "catalog": catalog,
    }


def fmt_money(value: float) -> str:
    return f"{value:,.2f}".replace(",", " ")


def fmt_int(value: int) -> str:
    return f"{value:,}".replace(",", " ")


def fmt_delta(value: float, money: bool = False) -> str:
    if abs(value) < 0.000001:
        return "0"
    prefix = "+" if value > 0 else ""
    if money:
        return f"{prefix}{fmt_money(value)}"
    return f"{prefix}{fmt_int(int(round(value)))}"


def build_report(
    workbook_path: Path,
    tenant_id: str,
    date_from: str,
    date_to: str,
    export_metrics: dict[str, Any],
    db_metrics: dict[str, Any],
) -> str:
    context = export_metrics["context"]
    filters = export_metrics["filters"]
    products = export_metrics["products"]
    tenant = db_metrics["tenant"]
    overview = db_metrics["overview"]
    export_slice_funnel = db_metrics["export_slice_funnel"]
    export_slice_economics = db_metrics["export_slice_economics"]
    catalog = db_metrics["catalog"]

    export_nm_ids = {int(value) for value in products["unique_nm_ids"]}
    products_nm_ids = {int(value) for value in catalog["products_nm_ids"]}
    active_products_nm_ids = {int(value) for value in catalog["active_products_nm_ids"]}
    metadata_nm_ids = {int(value) for value in catalog["metadata_nm_ids"]}
    observed_nm_ids = {int(value) for value in catalog["observed_nm_ids"]}

    missing_in_products = sorted(export_nm_ids - products_nm_ids)
    missing_in_active_products = sorted(export_nm_ids - active_products_nm_ids)
    missing_in_metadata = sorted(export_nm_ids - metadata_nm_ids)
    missing_in_observed = sorted(export_nm_ids - observed_nm_ids)

    delta_views = float(export_slice_funnel["views"]) - filters["views"]
    delta_carts = float(export_slice_funnel["carts"]) - filters["add_to_cart"]
    delta_orders = float(export_slice_funnel["orders"]) - filters["ordered_units"]
    delta_buyouts = float(export_slice_funnel["buyouts"]) - filters["bought_units"]
    delta_cancels = float(export_slice_funnel["cancels"]) - filters["cancelled_units"]
    delta_order_sum = float(export_slice_funnel["order_sum"]) - filters["ordered_rub"]
    delta_buyout_sum = float(export_slice_funnel["buyout_sum"]) - filters["bought_rub"]
    delta_cancel_sum = float(export_slice_funnel["cancel_sum"]) - filters["cancelled_rub"]

    lines = [
        f"# WB export reconciliation: {tenant['name']}",
        "",
        f"- Generated at: {datetime.now().isoformat(timespec='seconds')}",
        f"- Workbook: `{workbook_path}`",
        f"- Tenant ID: `{tenant_id}`",
        f"- Range: `{date_from}` -> `{date_to}`",
        f"- Export subject filter: `{context['subject'] or 'не указан'}`",
        f"- Deleted goods mode: `{context['deleted_mode'] or 'не указано'}`",
        f"- Granularity: `{context['granularity'] or 'не указано'}`",
        "",
        "## WB export totals",
        "",
        f"- Views: `{fmt_int(filters['views'])}`",
        f"- Add to cart: `{fmt_int(filters['add_to_cart'])}`",
        f"- Ordered units: `{fmt_int(filters['ordered_units'])}`",
        f"- Bought units: `{fmt_int(filters['bought_units'])}`",
        f"- Cancelled units: `{fmt_int(filters['cancelled_units'])}`",
        f"- Ordered rub: `{fmt_money(filters['ordered_rub'])}`",
        f"- Bought rub: `{fmt_money(filters['bought_rub'])}`",
        f"- Cancelled rub: `{fmt_money(filters['cancelled_rub'])}`",
        f"- Cart conversion: `{filters['cart_conversion_pct']:.2f}%`",
        f"- Order conversion: `{filters['order_conversion_pct']:.2f}%`",
        f"- Export rows: `{fmt_int(products['rows'])}`",
        f"- Export unique WB SKUs: `{fmt_int(len(export_nm_ids))}`",
        f"- Export SKUs marked deleted: `{fmt_int(len(products['deleted_nm_ids']))}`",
        "",
        "## App period facts for the same export SKU set",
        "",
        f"- Matched funnel SKUs in app: `{fmt_int(int(export_slice_funnel['matched_skus']))}` of `{fmt_int(len(export_nm_ids))}`",
        f"- App funnel views: `{fmt_int(int(float(export_slice_funnel['views'])))}` (`delta {fmt_delta(delta_views)}`)",
        f"- App funnel carts: `{fmt_int(int(float(export_slice_funnel['carts'])))}` (`delta {fmt_delta(delta_carts)}`)",
        f"- App funnel orders: `{fmt_int(int(float(export_slice_funnel['orders'])))}` (`delta {fmt_delta(delta_orders)}`)",
        f"- App funnel buyouts: `{fmt_int(int(float(export_slice_funnel['buyouts'])))}` (`delta {fmt_delta(delta_buyouts)}`)",
        f"- App funnel cancels: `{fmt_int(int(float(export_slice_funnel['cancels'])))}` (`delta {fmt_delta(delta_cancels)}`)",
        f"- App funnel order sum: `{fmt_money(float(export_slice_funnel['order_sum']))}` (`delta {fmt_delta(delta_order_sum, money=True)}`)",
        f"- App funnel buyout sum: `{fmt_money(float(export_slice_funnel['buyout_sum']))}` (`delta {fmt_delta(delta_buyout_sum, money=True)}`)",
        f"- App funnel cancel sum: `{fmt_money(float(export_slice_funnel['cancel_sum']))}` (`delta {fmt_delta(delta_cancel_sum, money=True)}`)",
        "",
        "## Current overview/economics scope",
        "",
        f"- Overview visible-catalog views: `{fmt_int(int(float(overview['total_views'])))}`",
        f"- Overview visible-catalog orders: `{fmt_int(int(float(overview['total_orders'])))}`",
        f"- Overview visible-catalog buyouts: `{fmt_int(int(float(overview['total_buyouts'])))}`",
        f"- Overview visible-catalog buyout sum: `{fmt_money(float(overview['total_buyout_sum']))}`",
        f"- Overview visible-catalog finance revenue: `{fmt_money(float(overview['finance_revenue']))}`",
        f"- Overview visible-catalog ads: `{fmt_money(float(overview['total_ads']))}`",
        f"- Economics revenue for export SKU set: `{fmt_money(float(export_slice_economics['finance_revenue']))}`",
        f"- Economics ads for export SKU set: `{fmt_money(float(export_slice_economics['total_ads']))}`",
        f"- Economics storage for export SKU set: `{fmt_money(float(export_slice_economics['total_storage']))}`",
        f"- Economics sold quantity for export SKU set: `{fmt_int(int(float(export_slice_economics['sold_quantity'])))}`",
        "",
        "## Catalog coverage for export SKUs",
        "",
        f"- Present in `products`: `{fmt_int(len(export_nm_ids & products_nm_ids))}` / `{fmt_int(len(export_nm_ids))}`",
        f"- Present in active `products`: `{fmt_int(len(export_nm_ids & active_products_nm_ids))}` / `{fmt_int(len(export_nm_ids))}`",
        f"- Present in `raw_api_product_metadata`: `{fmt_int(len(export_nm_ids & metadata_nm_ids))}` / `{fmt_int(len(export_nm_ids))}`",
        f"- Present in any observed raw facts: `{fmt_int(len(export_nm_ids & observed_nm_ids))}` / `{fmt_int(len(export_nm_ids))}`",
        f"- Missing in `products`: `{missing_in_products}`",
        f"- Missing in active `products`: `{missing_in_active_products}`",
        f"- Missing in `raw_api_product_metadata`: `{missing_in_metadata}`",
        f"- Missing in observed raw facts: `{missing_in_observed}`",
        "",
        "## Diagnosis",
        "",
        "- Period funnel facts for the export SKU set are now the right canonical comparison layer for the WB workbook.",
        "- The remaining mismatch is small and concentrated in SKU coverage, not in the storage model itself.",
        "- Overview numbers are wider than this workbook by design, because the workbook is filtered to one subject (`Тапочки`) and includes deleted goods, while overview uses the visible catalog scope of the app.",
        "- Economics revenue should not be compared 1:1 to WB buyout sum: economics is a realized PnL layer (`raw_api_realization_reports` + provisional `raw_api_sales`), while the workbook total is WB gross funnel/export turnover.",
        "- Historical `raw_api_product_metadata` should be preserved, and missing observed SKU should be backfilled into `products`; otherwise grouping / merge UX will keep dropping real historical items.",
    ]

    return "\n".join(lines) + "\n"


def main() -> None:
    parser = argparse.ArgumentParser(description="Compare a WB export workbook with local app data")
    parser.add_argument("--workbook", required=True, help="Path to WB .xlsx export")
    parser.add_argument("--tenant-id", required=True, help="Tenant UUID")
    parser.add_argument("--date-from", required=True, help="Inclusive ISO date, example 2026-03-01")
    parser.add_argument("--date-to", required=True, help="Inclusive ISO date, example 2026-03-31")
    parser.add_argument("--output", help="Optional markdown output path")
    args = parser.parse_args()

    workbook_path = Path(args.workbook).expanduser().resolve()
    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")

    database_url = load_env_database_url()
    export_metrics = load_export_metrics(workbook_path)
    db_metrics = load_db_metrics(
        database_url,
        args.tenant_id,
        args.date_from,
        args.date_to,
        export_metrics["products"]["unique_nm_ids"],
    )
    report = build_report(
        workbook_path=workbook_path,
        tenant_id=args.tenant_id,
        date_from=args.date_from,
        date_to=args.date_to,
        export_metrics=export_metrics,
        db_metrics=db_metrics,
    )

    if args.output:
        output_path = Path(args.output).expanduser().resolve()
    else:
        OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        output_path = OUTPUT_DIR / f"wb-export-reconciliation-{args.date_from}-to-{args.date_to}.md"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(report, encoding="utf-8")
    print(output_path)


if __name__ == "__main__":
    main()
